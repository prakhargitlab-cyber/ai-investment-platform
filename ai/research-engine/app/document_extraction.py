"""Provider-neutral document/file text extraction boundary.

This module exists to separate DOCUMENT/FILE EXTRACTION from everything
downstream of it (evidence interpretation, candidate drafts, canonical
persistence, Research Readiness). Today only
``app.manual_evidence.ManualEvidenceIngestor`` calls it, but it is
deliberately independent of that caller so the platform's existing
``LlmProvider`` abstraction -- and, later, a future Ollama/local-model
provider built on it -- can reuse the SAME extraction step: an LLM-based
evidence interpreter would consume the same ``ExtractedDocument`` this
module produces, never re-implement file parsing itself.

Hard boundary (do not import any of these here, or add an import that
creates a dependency on them):
    - app.research_readiness / app.research_readiness_runtime
    - app.manual_evidence (the relationship is one-directional: that module
      imports this one, never the reverse)
    - anything Equity Radar / ETF Radar specific
    - Ollama, or any specific LLM provider
This module knows nothing about evidence types (SHAREHOLDING, CURRENT_NEWS,
...), readiness statuses, or scoring. It turns bytes into text/pages/
tables; what that text MEANS is entirely the job of the evidence
interpretation layer (see app.evidence_interpretation), which consumes
this module's output as plain data.

Extraction method implementations are REUSED, not duplicated: PNG/JPEG and
scanned-PDF OCR delegate to the already-independent, already-bounded
app.image_evidence_extraction (subprocess-isolated tesseract/pdftoppm,
with its own file-size/pixel/page/timeout bounds -- see that module's
docstring for the full resource-safety rationale, unchanged by this
refactor). Native PDF/DOCX/CSV/TXT extraction uses pypdf/python-docx
directly, exactly as app.manual_evidence did before this refactor -- moved
here, not rewritten.

MIME handling: ``DocumentInput.mime_type`` must already be detected/
validated by the caller (see app.manual_evidence.detect_mime_type). This
module never re-sniffs it from bytes or filename -- it is handed through
cleanly and dispatched on directly, so MIME detection has exactly one
owner in the codebase.
"""
from __future__ import annotations

import hashlib
import io
import logging
from dataclasses import dataclass, field
from typing import Any

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Supported MIME types
# ---------------------------------------------------------------------------

TXT_MIME = "text/plain"
CSV_MIME = "text/csv"
PDF_MIME = "application/pdf"
PNG_MIME = "image/png"
JPEG_MIME = "image/jpeg"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# Every MIME type this extractor has *an extraction path for*. Whether DOCX
# extraction is genuinely usable right now additionally depends on
# docx_extraction_available() (python-docx being importable) -- mirrors the
# existing OCR fail-closed pattern in app.image_evidence_extraction, now
# generalized to any format with an optional runtime dependency.
EXTRACTABLE_MIME_TYPES: tuple[str, ...] = (TXT_MIME, CSV_MIME, PDF_MIME, PNG_MIME, JPEG_MIME, DOCX_MIME)


def docx_extraction_available() -> bool:
    """Lightweight (no-parse) runtime capability check: is the python-docx
    package importable right now? Mirrors
    app.image_evidence_extraction.image_ocr_available()'s fail-closed
    pattern -- a deployment missing the dependency must truthfully stop
    advertising DOCX support rather than accept uploads and fail later."""
    try:
        import docx  # noqa: F401
    except ImportError:
        return False
    return True


# ---------------------------------------------------------------------------
# Input / output contracts
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class DocumentInput:
    """What DocumentExtractor needs to extract text from a file.

    ``mime_type`` must already be detected/validated by the caller and is
    passed through verbatim -- seeApp module docstring's MIME handling note.
    ``size_bytes`` defaults to ``len(file_bytes)`` when not supplied;
    callers that already know the size (e.g. from an HTTP Content-Length)
    may pass it explicitly to avoid re-measuring large payloads twice.
    """
    file_bytes: bytes
    mime_type: str
    filename: str | None = None
    size_bytes: int | None = None
    # Optional caller-supplied label/URL for the evidence item this file
    # belongs to. Never required, never interpreted by this module --
    # purely carried through into ExtractedDocument.provenance for the
    # caller's own audit trail.
    source_reference: str | None = None

    @property
    def byte_size(self) -> int:
        return self.size_bytes if self.size_bytes is not None else len(self.file_bytes)


@dataclass
class ExtractedDocument:
    """Normalized extraction output. Provider/evidence-type neutral --
    nothing here is specific to any Research Readiness requirement.

    ``confidence`` is populated ONLY when a specific extraction step
    genuinely reports one; it is None for every deterministic native
    extraction (TXT/CSV/PDF-text/DOCX), which has no meaningful notion of
    confidence. OCR extraction (image/scanned-PDF) also leaves this None
    today -- tesseract's own per-word confidence is not threaded through,
    so inventing a document-level number here would be exactly the
    fabricated-confidence the task prohibits.
    """
    text: str
    pages: list[str] | None = None
    tables: list[list[list[str]]] | None = None
    metadata: dict[str, Any] = field(default_factory=dict)
    extraction_method: str = ""
    warnings: list[str] = field(default_factory=list)
    provenance: dict[str, Any] = field(default_factory=dict)
    confidence: float | None = None


class DocumentExtractionError(Exception):
    """Raised ONLY for a genuinely unsupported MIME type passed to
    DocumentExtractor.extract(). Every other failure mode (corrupt file,
    OCR unavailable, subprocess timeout, oversized input, empty file) is
    reported truthfully inside the returned ExtractedDocument's
    warnings/metadata instead of being raised -- a caller building a
    review draft should never need a try/except around a routine
    extraction failure to get a usable (if empty) draft back."""


def _base_provenance(document: DocumentInput) -> dict[str, Any]:
    return {
        "contentHash": hashlib.sha256(document.file_bytes).hexdigest(),
        "mimeType": document.mime_type,
        "filename": document.filename,
        "sizeBytes": document.byte_size,
        "sourceReference": document.source_reference,
    }


class DocumentExtractor:
    """``extract(DocumentInput) -> ExtractedDocument``.

    Dispatches purely on ``document.mime_type`` (never re-sniffed here).
    Callable and testable entirely on its own -- it does not import
    ``ResearchReadinessRuntime`` or anything evidence-type-specific, which
    is exactly what lets a future LLM-based interpreter reuse this same
    extraction step for a document an evidence interpreter has not been
    written for yet.
    """

    def extract(self, document: DocumentInput) -> ExtractedDocument:
        mime = document.mime_type
        if mime == TXT_MIME:
            return self._extract_plain_text(document)
        if mime == CSV_MIME:
            return self._extract_csv(document)
        if mime == PDF_MIME:
            return self._extract_pdf(document)
        if mime in (PNG_MIME, JPEG_MIME):
            return self._extract_image(document)
        if mime == DOCX_MIME:
            return self._extract_docx(document)
        raise DocumentExtractionError(f"UNSUPPORTED_MIME_TYPE: {mime}")

    # -- native text formats --------------------------------------------

    def _extract_plain_text(self, document: DocumentInput) -> ExtractedDocument:
        text = document.file_bytes.decode("utf-8", errors="replace")
        return ExtractedDocument(
            text=text,
            extraction_method="TXT_NATIVE",
            provenance=_base_provenance(document),
        )

    def _extract_csv(self, document: DocumentInput) -> ExtractedDocument:
        import csv as csv_module

        text = document.file_bytes.decode("utf-8", errors="replace")
        tables: list[list[list[str]]] | None = None
        warnings: list[str] = []
        try:
            rows = list(csv_module.reader(io.StringIO(text)))
            if rows:
                tables = [rows]
        except csv_module.Error as exc:
            warnings.append(f"CSV_ROW_PARSE_FAILED: {exc}")
        return ExtractedDocument(
            text=text,
            tables=tables,
            extraction_method="CSV_NATIVE",
            warnings=warnings,
            provenance=_base_provenance(document),
        )

    # -- PDF (native text first, bounded OCR fallback for scanned PDFs) --

    def _extract_pdf(self, document: DocumentInput) -> ExtractedDocument:
        provenance = _base_provenance(document)
        try:
            from pypdf import PdfReader
        except ImportError:
            return ExtractedDocument(
                text="", extraction_method="PDF_UNAVAILABLE",
                warnings=["PDF_EXTRACTION_UNAVAILABLE: pypdf is not installed"],
                provenance=provenance,
            )
        try:
            reader = PdfReader(io.BytesIO(document.file_bytes))
            pages = [page.extract_text() or "" for page in reader.pages]
        except Exception as exc:  # noqa: BLE001 -- any decode failure is a truthful rejection
            logger.warning("pdf_text_extraction_failed reason=%s", exc)
            return ExtractedDocument(
                text="", extraction_method="PDF_INVALID",
                warnings=[f"INVALID_PDF: {exc}"],
                provenance=provenance,
            )
        native_text = "\n".join(pages)
        if native_text.strip():
            return ExtractedDocument(
                text=native_text, pages=pages,
                extraction_method="PDF_TEXT_NATIVE",
                provenance=provenance,
            )
        # No extractable text layer at all -- fall back to bounded scanned-
        # PDF OCR (never attempted when native text already exists, so a
        # text PDF is never OCR'd unnecessarily).
        return self._extract_scanned_pdf(document, provenance)

    def _extract_scanned_pdf(self, document: DocumentInput, provenance: dict[str, Any]) -> ExtractedDocument:
        from app.image_evidence_extraction import extract_text_from_scanned_pdf

        result = extract_text_from_scanned_pdf(document.file_bytes)
        metadata = {"ocrStatus": result.status, "pageCount": result.page_count}
        if result.status != "EXTRACTED":
            logger.info("scanned_pdf_ocr_not_extracted status=%s reason=%s", result.status, result.reason)
            return ExtractedDocument(
                text="", extraction_method="PDF_SCANNED_OCR_NOT_EXTRACTED",
                warnings=[result.reason] if result.reason else [],
                metadata=metadata,
                provenance=provenance,
            )
        return ExtractedDocument(
            text=result.text or "", extraction_method="PDF_SCANNED_OCR",
            metadata=metadata,
            provenance=provenance,
        )

    # -- images (bounded, subprocess-isolated OCR) -----------------------

    def _extract_image(self, document: DocumentInput) -> ExtractedDocument:
        from app.image_evidence_extraction import extract_text_from_image

        provenance = _base_provenance(document)
        result = extract_text_from_image(document.file_bytes, document.mime_type, filename=document.filename)
        metadata = {"ocrStatus": result.status}
        if result.status != "EXTRACTED":
            logger.info("image_ocr_not_extracted status=%s reason=%s", result.status, result.reason)
            return ExtractedDocument(
                text="", extraction_method="IMAGE_OCR_NOT_EXTRACTED",
                warnings=[result.reason] if result.reason else [],
                metadata=metadata,
                provenance=provenance,
            )
        return ExtractedDocument(
            text=result.text or "", extraction_method="IMAGE_OCR",
            metadata=metadata,
            provenance=provenance,
        )

    # -- DOCX (paragraphs + tables, no LibreOffice/Word runtime) ---------

    def _extract_docx(self, document: DocumentInput) -> ExtractedDocument:
        provenance = _base_provenance(document)
        if not docx_extraction_available():
            return ExtractedDocument(
                text="", extraction_method="DOCX_UNAVAILABLE",
                warnings=["DOCX_EXTRACTION_UNAVAILABLE: python-docx is not installed"],
                provenance=provenance,
            )
        try:
            from docx import Document as DocxDocument
        except ImportError:
            return ExtractedDocument(
                text="", extraction_method="DOCX_UNAVAILABLE",
                warnings=["DOCX_EXTRACTION_UNAVAILABLE: python-docx is not installed"],
                provenance=provenance,
            )
        try:
            doc = DocxDocument(io.BytesIO(document.file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text]
            tables: list[list[list[str]]] = [
                [[cell.text for cell in row.cells] for row in table.rows]
                for table in doc.tables
            ]
            core_props = doc.core_properties
            metadata = {
                "title": core_props.title or None,
                "author": core_props.author or None,
                "created": core_props.created.isoformat() if core_props.created else None,
                "tableCount": len(tables),
                "paragraphCount": len(paragraphs),
            }
        except Exception as exc:  # noqa: BLE001 -- any decode failure is a truthful rejection
            logger.warning("docx_text_extraction_failed reason=%s", exc)
            return ExtractedDocument(
                text="", extraction_method="DOCX_INVALID",
                warnings=[f"INVALID_DOCX: {exc}"],
                provenance=provenance,
            )
        table_text = "\n".join(
            "\n".join(" | ".join(cell for cell in row) for row in table)
            for table in tables
        )
        text = "\n".join(part for part in (*paragraphs, table_text) if part)
        return ExtractedDocument(
            text=text,
            tables=tables or None,
            metadata=metadata,
            extraction_method="DOCX_NATIVE",
            provenance=provenance,
        )
